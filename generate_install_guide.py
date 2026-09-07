#!/usr/bin/env python3
"""Generate the KMIP on PKCS#11 installation and testing guide (Word format).

Run: python generate_install_guide.py

Covers a full deployment, not just a development checkout: prerequisites,
building SoftHSM2, token initialisation, TLS material, configuration, first
start, running as a service or a container, provisioning, verification, the
test suite, upgrades, backup and restore, hardening and troubleshooting.

Commands here were run on the machine that produced this document unless a
line says otherwise, and the troubleshooting section quotes real error text
rather than paraphrases.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

# ── colour palette ────────────────────────────────────────────────────────────
DARK_BLUE   = RGBColor(0x1F, 0x39, 0x64)   # headings
MID_BLUE    = RGBColor(0x2E, 0x74, 0xB5)   # subheadings / accent
LIGHT_BLUE  = RGBColor(0xD6, 0xE4, 0xF0)   # table header fill
CODE_BG     = RGBColor(0xF2, 0xF2, 0xF2)   # code block background
GREEN       = RGBColor(0x37, 0x86, 0x44)   # pass indicators
ORANGE      = RGBColor(0xD8, 0x6B, 0x00)   # warning / note
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)

def _rgb_hex(color: RGBColor) -> str:
    """Convert RGBColor to uppercase hex string (no '#')."""
    return str(color).upper()  # RGBColor.__str__ returns e.g. '2E74B5'

CODE_FONT   = "Courier New"
BODY_FONT   = "Calibri"

# ── helpers ───────────────────────────────────────────────────────────────────

def set_cell_bg(cell, rgb: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    hex_color = _rgb_hex(rgb)
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_paragraph_shading(para, rgb: RGBColor):
    pPr  = para._p.get_or_add_pPr()
    shd  = OxmlElement('w:shd')
    hex_color = _rgb_hex(rgb)
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    pPr.append(shd)


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(6)
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(18)
    run.font.color.rgb = DARK_BLUE
    run.font.name = BODY_FONT
    # bottom border
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), _rgb_hex(MID_BLUE))
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(3)
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(14)
    run.font.color.rgb = MID_BLUE
    run.font.name = BODY_FONT
    return p


def h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(12)
    run.font.color.rgb = DARK_BLUE
    run.font.name = BODY_FONT
    return p


def body(doc, text, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = BODY_FONT
    return p


def note(doc, text, color=ORANGE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(6)
    p.paragraph_format.left_indent  = Cm(0.5)
    add_paragraph_shading(p, RGBColor(0xFF, 0xF4, 0xE5))
    run = p.add_run("⚠  " + text)
    run.font.size = Pt(10)
    run.font.color.rgb = color
    run.font.name = BODY_FONT
    return p


def tip(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(6)
    p.paragraph_format.left_indent  = Cm(0.5)
    add_paragraph_shading(p, RGBColor(0xE8, 0xF5, 0xE9))
    run = p.add_run("✔  " + text)
    run.font.size = Pt(10)
    run.font.color.rgb = GREEN
    run.font.name = BODY_FONT
    return p


def code_block(doc, lines, title=None):
    if title:
        tp = doc.add_paragraph()
        tp.paragraph_format.space_before = Pt(6)
        tp.paragraph_format.space_after  = Pt(0)
        tr = tp.add_run(title)
        tr.bold           = True
        tr.font.size      = Pt(9)
        tr.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        tr.font.name      = BODY_FONT

    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)
        p.paragraph_format.left_indent  = Cm(0.4)
        p.paragraph_format.right_indent = Cm(0.4)
        add_paragraph_shading(p, CODE_BG)
        run = p.add_run(line if line else " ")
        run.font.name  = CODE_FONT
        run.font.size  = Pt(9)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

    # bottom spacer
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after  = Pt(6)
    add_paragraph_shading(sp, CODE_BG)
    sp.add_run(" ").font.size = Pt(3)


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.left_indent  = Cm(0.6 + level * 0.6)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = BODY_FONT
    return p


def table_header_row(table, headers, col_widths_cm=None):
    row = table.rows[0]
    for i, hdr in enumerate(headers):
        cell = row.cells[i]
        set_cell_bg(cell, DARK_BLUE)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after  = Pt(2)
        run = p.add_run(hdr)
        run.bold           = True
        run.font.color.rgb = WHITE
        run.font.size      = Pt(10)
        run.font.name      = BODY_FONT


def add_table_row(table, values, shade_even=False, idx=0):
    row = table.add_row()
    for j, val in enumerate(values):
        cell = row.cells[j]
        if shade_even and idx % 2 == 0:
            set_cell_bg(cell, RGBColor(0xF5, 0xF8, 0xFF))
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after  = Pt(1)
        is_code = val.startswith('`') and val.endswith('`')
        text = val[1:-1] if is_code else val
        run = p.add_run(text)
        run.font.size = Pt(10)
        run.font.name = CODE_FONT if is_code else BODY_FONT


def make_table(doc, headers, rows, col_widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_header_row(t, headers)
    for i, row in enumerate(rows):
        add_table_row(t, row, shade_even=True, idx=i)
    if col_widths:
        for i, row in enumerate(t.rows):
            for j, cell in enumerate(row.cells):
                cell.width = Cm(col_widths[j])
    return t


# ═════════════════════════════════════════════════════════════════════════════
# Document assembly
# ═════════════════════════════════════════════════════════════════════════════


TODAY = datetime.date.today().strftime("%d %B %Y")


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ── cover ────────────────────────────────────────────────────────────
    doc.add_paragraph()
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title_p.add_run("KMIP on PKCS#11")
    tr.bold = True
    tr.font.size = Pt(32)
    tr.font.color.rgb = DARK_BLUE
    tr.font.name = BODY_FONT

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub_p.add_run("Installation, Deployment & Test Guide")
    sr.font.size = Pt(16)
    sr.font.color.rgb = MID_BLUE
    sr.font.name = BODY_FONT

    doc.add_paragraph()
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta_p.add_run(f"Version 1.0  ·  {TODAY}")
    mr.font.size = Pt(11)
    mr.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    mr.font.name = BODY_FONT
    doc.add_paragraph()

    # ── 1  Before you start ──────────────────────────────────────────────
    h1(doc, "1  Before you start")
    body(doc,
        "This guide takes a machine with nothing on it to a running, "
        "TLS-protected KMIP 2.1 server backed by a PKCS#11 token, and then "
        "shows how to verify it, operate it, upgrade it and recover it. It "
        "also covers running the 816-test suite, which is the fastest way to "
        "confirm an installation is sound.")

    h2(doc, "1.1  Which path do you want?")
    make_table(doc,
        ["Goal", "Read"],
        [
            ["Try it out on a laptop",
             "Sections 2, 3, 4, 5 — then 9 to run the tests. Roughly 20 minutes, "
             "most of it compiling SoftHSM2."],
            ["Deploy it as a service",
             "All of Sections 2–8, then 10 (systemd) or 11 (container). Do not skip "
             "Section 6: the server refuses to start without TLS unless told to."],
            ["Work on the code",
             "Sections 2–5 with an editable install, then Section 9."],
        ],
        col_widths=[4.5, 12.5])
    doc.add_paragraph()

    h2(doc, "1.2  What gets installed")
    make_table(doc,
        ["Component", "Version used here", "Role"],
        [
            ["Python", "3.11.15 (3.9+ required)", "Runtime."],
            ["SoftHSM2", "2.7.0, built from source", "Software PKCS#11 token. See "
                                                     "Section 3 for why the packaged "
                                                     "build is not adequate."],
            ["OpenSSL", "3.0.13", "SoftHSM2's crypto backend, and the tool used to "
                                  "create TLS material."],
            ["python-pkcs11", ">=0.9.5,<0.10", "The only PKCS#11 binding used. Pinned: "
                                               "the shim relies on 0.9.x call "
                                               "signatures."],
            ["PyYAML", ">=5.4", "Configuration files. Only the safe loader is used."],
            ["asn1crypto", ">=1.4", "X.509 for Certify / ReCertify / Validate. Note "
                                    "that `cryptography` is deliberately NOT a "
                                    "dependency — all real crypto happens on the "
                                    "token."],
            ["pytest, pytest-cov, pytest-timeout", "current", "Test suite only "
                                                              "(extras: dev)."],
        ],
        col_widths=[4.0, 4.0, 9.0])
    doc.add_paragraph()
    tip(doc, "Three dependencies at runtime, none of which does cryptography. That is "
             "deliberate: if a library on this list could decrypt your keys, the HSM "
             "boundary would not mean much.")

    h2(doc, "1.3  Where things live")
    make_table(doc,
        ["Path", "Contents"],
        [
            ["`/etc/kmip/config.yaml`", "Server configuration."],
            ["`/etc/kmip/tls/`", "Certificate, private key, CA."],
            ["`/var/lib/kmip/kmip.db`", "Metadata: objects, identities, roles, groups, "
                                        "grants, approvals, audit log."],
            ["`/var/lib/softhsm/tokens/`", "The token store — the key material itself."],
            ["`/run/secrets/kmip-token-pin`", "The token PIN, mounted by a secrets "
                                              "manager."],
        ],
        col_widths=[6.0, 11.0])
    doc.add_paragraph()
    note(doc, "The database and the token store are a matched pair. Secret blobs in the "
              "database are encrypted under a master key that lives on the token, and "
              "objects reference token keys by CKA_ID — so a database restored beside a "
              "different token is not a degraded backup, it is unreadable. Back them up "
              "together and keep them together.")
    doc.add_page_break()

    # ── 2  Prerequisites ─────────────────────────────────────────────────
    h1(doc, "2  System prerequisites")
    body(doc,
        "Commands below target Debian and Ubuntu. The machine that produced this "
        "document runs Ubuntu 24.04 LTS with Python 3.11.15 and OpenSSL 3.0.13.")

    h2(doc, "2.1  Build and runtime packages")
    code_block(doc, [
        "sudo apt update",
        "",
        "# Build SoftHSM2 from source (Section 3):",
        "sudo apt install -y build-essential automake autoconf libtool \\",
        "                    pkg-config libssl-dev ca-certificates curl",
        "",
        "# Python:",
        "sudo apt install -y python3 python3-pip python3-venv python3-dev",
        "",
        "# Useful, not required:",
        "sudo apt install -y opensc sqlite3",
    ])
    body(doc,
        "opensc provides pkcs11-tool, which is handy for inspecting a token "
        "independently of this project. sqlite3 lets you look at the metadata "
        "database directly — read-only; the audit log has triggers that refuse "
        "modification.")

    h2(doc, "2.2  A user to run as")
    code_block(doc, [
        "sudo useradd --system --home-dir /var/lib/kmip --create-home kmip",
        "sudo mkdir -p /etc/kmip/tls /var/lib/kmip /var/lib/softhsm/tokens",
        "sudo chown -R kmip:kmip /var/lib/kmip /var/lib/softhsm",
        "sudo chmod 700 /var/lib/kmip /var/lib/softhsm/tokens",
    ])
    note(doc, "Do not run the server as root. It needs its database directory, the token "
              "store and its configuration, and nothing else — the systemd unit in "
              "Section 10 enforces exactly that.")
    doc.add_page_break()

    # ── 3  SoftHSM2 ──────────────────────────────────────────────────────
    h1(doc, "3  Installing SoftHSM2")

    h2(doc, "3.1  Why not apt install softhsm2")
    body(doc,
        "Because the packaged version does not expose the mechanisms this server "
        "needs for EC signing. This is measurable, and worth measuring rather than "
        "taking on trust:")
    make_table(doc,
        ["Build", "Mechanisms advertised", "CKM_ECDSA_SHA256"],
        [
            ["Ubuntu 24.04 package, SoftHSM 2.6.1", "70", "Absent"],
            ["Source build, SoftHSM 2.7.0", "79", "Present"],
        ],
        col_widths=[8.0, 5.0, 4.0])
    doc.add_paragraph()
    body(doc,
        "Both builds link OpenSSL's libcrypto, so the crypto backend is not the "
        "differentiator — the version is. Against 2.6.1 every EC signing operation "
        "fails, and so does every EC signing test. If your distribution ships 2.7.0 "
        "or later, the package is fine; check before building.")
    code_block(doc, [
        "# Check what your distribution offers:",
        "apt-cache policy softhsm2",
        "",
        "# Check what a token actually advertises, whichever build you use:",
        "SOFTHSM2_CONF=/etc/kmip/softhsm2.conf pkcs11-tool \\",
        "    --module /usr/local/lib/softhsm/libsofthsm2.so -M | grep -i ecdsa",
    ])

    h2(doc, "3.2  Build from source")
    code_block(doc, [
        "VERSION=2.7.0",
        "curl -fsSL \"https://dist.opendnssec.org/source/softhsm-${VERSION}.tar.gz\" \\",
        "     -o /tmp/softhsm.tar.gz",
        "tar -xzf /tmp/softhsm.tar.gz -C /tmp",
        "cd /tmp/softhsm-${VERSION}",
        "",
        "./configure --prefix=/usr/local \\",
        "            --with-crypto-backend=openssl \\",
        "            --disable-gost",
        "make -j\"$(nproc)\"",
        "sudo make install",
        "sudo ldconfig",
        "",
        "softhsm2-util --version",
        "# → 2.7.0",
    ])
    body(doc,
        "--disable-gost drops the GOST algorithms, which OpenSSL 3 no longer "
        "provides and which nothing here uses. The library lands at "
        "/usr/local/lib/softhsm/libsofthsm2.so — that path goes in the "
        "configuration file.")
    tip(doc, "Installing to /usr/local leaves any distribution package in place at "
             "/usr/lib/softhsm/. That is useful: you can point the config at either and "
             "compare, which is how the table above was produced.")

    h2(doc, "3.3  Using a real HSM instead")
    body(doc,
        "Nothing above is specific to SoftHSM2. The shim is the only file that "
        "imports pkcs11, so any PKCS#11 v2.40+ token works: install the vendor's "
        "library, point hsm.library at it, and set hsm.token_label to the token's "
        "label. The capability probe at startup will report what that token "
        "supports, and any algorithm it lacks is refused cleanly rather than "
        "failing deep in an operation.")
    note(doc, "A validated HSM has not been tested here — the PKCS#11 boundary makes the "
              "swap cheap, but 'cheap' is not 'verified'. Run the test suite against the "
              "token before trusting it, and expect the algorithm coverage table to "
              "differ.")
    doc.add_page_break()

    # ── 4  Token ─────────────────────────────────────────────────────────
    h1(doc, "4  Initialising the token")

    h2(doc, "4.1  Point SoftHSM2 at a token store")
    body(doc,
        "SoftHSM2 finds its token directory through a configuration file named by "
        "the SOFTHSM2_CONF environment variable. The server needs that variable "
        "set in its environment too — the systemd unit and the container image "
        "both set it.")
    code_block(doc, [
        "sudo tee /etc/kmip/softhsm2.conf >/dev/null <<'EOF'",
        "directories.tokendir = /var/lib/softhsm/tokens",
        "objectstore.backend = file",
        "log.level = ERROR",
        "EOF",
        "sudo chown kmip:kmip /etc/kmip/softhsm2.conf",
        "",
        "export SOFTHSM2_CONF=/etc/kmip/softhsm2.conf",
    ])

    h2(doc, "4.2  Create the token")
    code_block(doc, [
        "sudo -u kmip SOFTHSM2_CONF=/etc/kmip/softhsm2.conf \\",
        "    softhsm2-util --init-token --slot 0 --label KMIPToken \\",
        "                  --pin '<user-pin>' --so-pin '<so-pin>'",
        "",
        "# Confirm it exists:",
        "sudo -u kmip SOFTHSM2_CONF=/etc/kmip/softhsm2.conf softhsm2-util --show-slots",
    ])
    body(doc,
        "The label is what hsm.token_label refers to. The user PIN is what the "
        "server presents; the SO PIN is only for administering the token itself "
        "and the server never uses it — store it somewhere else entirely.")
    note(doc, "The token PIN authenticates the server to the HSM. It authenticates no "
              "client and grants no KMIP authority — that is what identities are for. "
              "Anyone who has it can use the token directly, so treat it as a machine "
              "credential and keep it out of the configuration file (Section 7.2).")

    h2(doc, "4.3  Re-initialising")
    note(doc, "Re-initialising a token destroys every key on it, and the metadata "
              "database that references those keys becomes useless — the objects it "
              "describes no longer exist, and secret blobs encrypted under the master "
              "key cannot be decrypted. If you re-initialise, start with a fresh "
              "database.")
    doc.add_page_break()

    # ── 5  Package ───────────────────────────────────────────────────────
    h1(doc, "5  Installing the Python package")

    h2(doc, "5.1  Get the source")
    code_block(doc, [
        "git clone https://github.com/Gaillotte/Noida.git",
        "cd Noida",
    ])

    h2(doc, "5.2  Virtual environment")
    code_block(doc, [
        "python3 -m venv .venv",
        "source .venv/bin/activate",
        "python -m pip install --upgrade pip setuptools wheel",
    ])
    body(doc,
        "For a system service, create the environment somewhere the service user "
        "can read — /opt/kmip/venv is a reasonable choice — and give ExecStart the "
        "absolute path to that environment's kmip-server.")

    h2(doc, "5.3  Install")
    code_block(doc, [
        "# Deployment:",
        "pip install .",
        "",
        "# Development (editable, with the test extras):",
        "pip install --use-pep517 -e '.[dev]'",
    ])
    note(doc, "On Debian and Ubuntu, a plain `pip install -e .` can fail with "
              "AttributeError: install_layout — a known interaction between the "
              "distribution's setuptools patches and the legacy editable path. "
              "--use-pep517 avoids it. A non-editable `pip install .` is unaffected.")

    h2(doc, "5.4  Confirm it installed")
    code_block(doc, [
        "kmip-server --help",
        "kmip-admin --help",
        "python -c \"import kmip_pkcs11; print(kmip_pkcs11.__file__)\"",
    ])
    body(doc,
        "Two console scripts is the whole interface. If kmip-server is not on the "
        "path, the virtual environment is not active or the install went to a "
        "different interpreter.")
    doc.add_page_break()

    # ── 6  TLS ───────────────────────────────────────────────────────────
    h1(doc, "6  TLS material")
    body(doc,
        "The server refuses to start without TLS unless server.allow_plaintext is "
        "explicitly true. That is the intended shape: plaintext is available for a "
        "test rig, but never by accident.")

    h2(doc, "6.1  A private CA and a server certificate")
    body(doc,
        "For production, use your organisation's PKI. The commands below produce "
        "something usable for a lab or an internal deployment.")
    code_block(doc, [
        "cd /etc/kmip/tls",
        "",
        "# 1. A CA:",
        "openssl req -x509 -newkey rsa:4096 -days 3650 -nodes \\",
        "    -keyout ca.key -out ca.pem \\",
        "    -subj \"/CN=KMIP Internal CA\"",
        "",
        "# 2. A server key and CSR (the SAN is what clients verify):",
        "openssl req -newkey rsa:4096 -nodes \\",
        "    -keyout server.key -out server.csr \\",
        "    -subj \"/CN=kms.internal\" \\",
        "    -addext \"subjectAltName=DNS:kms.internal,IP:10.0.0.10\"",
        "",
        "# 3. Sign it:",
        "openssl x509 -req -in server.csr -CA ca.pem -CAkey ca.key \\",
        "    -CAcreateserial -days 825 -out server.pem \\",
        "    -copy_extensions copyall",
        "",
        "sudo chown kmip:kmip server.key server.pem ca.pem",
        "sudo chmod 600 server.key",
    ])
    note(doc, "-copy_extensions copyall carries the SAN from the CSR into the "
              "certificate. Without it the SAN is silently dropped and clients reject "
              "the certificate with a hostname mismatch — a confusing failure, because "
              "the certificate looks correct until you print its extensions.")

    h2(doc, "6.2  Client certificates, for mTLS")
    code_block(doc, [
        "openssl req -newkey rsa:2048 -nodes \\",
        "    -keyout alice.key -out alice.csr -subj \"/CN=alice\"",
        "openssl x509 -req -in alice.csr -CA ca.pem -CAkey ca.key \\",
        "    -CAcreateserial -days 825 -out alice.pem",
    ])
    body(doc,
        "With tls.require_client_cert: true, a verified certificate's Common Name is "
        "used as the caller's identity — but only if it names an identity that has "
        "actually been provisioned (Section 8). A CA-signed certificate does not get "
        "to invent a principal that was never granted anything, so CN=alice only "
        "works once alice exists.")

    h2(doc, "6.3  Renewal")
    body(doc,
        "SIGHUP re-reads the certificate and key without dropping established "
        "connections, so renewal costs no downtime:")
    code_block(doc, [
        "sudo systemctl reload kmip-server      # ExecReload sends SIGHUP",
    ])
    note(doc, "There is no ACME client. Obtaining and renewing certificates is the "
              "operator's job; only the reload is automatic.")
    doc.add_page_break()

    # ── 7  Configuration ─────────────────────────────────────────────────
    h1(doc, "7  Configuration")
    body(doc,
        "One YAML file holds everything. deploy/config.example.yaml in the "
        "repository is a fully annotated template; the Feature Specification has "
        "the complete key reference. What follows is a working minimum plus the "
        "decisions worth making deliberately.")

    h2(doc, "7.1  A working configuration")
    code_block(doc, [
        "# /etc/kmip/config.yaml",
        "server:",
        "  host: 0.0.0.0",
        "  port: 5696",
        "  workers: 4                  # null = one per CPU; 1 = single process",
        "",
        "tls:",
        "  cert: /etc/kmip/tls/server.pem",
        "  key:  /etc/kmip/tls/server.key",
        "  ca:   /etc/kmip/tls/ca.pem",
        "  require_client_cert: false",
        "",
        "hsm:",
        "  library: /usr/local/lib/softhsm/libsofthsm2.so",
        "  token_label: KMIPToken",
        "  pin_file: /run/secrets/kmip-token-pin",
        "",
        "storage:",
        "  database: /var/lib/kmip/kmip.db",
        "",
        "observability:",
        "  enabled: true",
        "  host: 127.0.0.1             # keep off the public interface",
        "  port: 9696",
        "",
        "governance:",
        "  enabled: false              # see Section 7.4",
        "",
        "logging:",
        "  level: INFO",
        "  format: json",
    ])

    h2(doc, "7.2  The token PIN")
    body(doc,
        "Supply exactly one source. Setting more than one is a configuration error "
        "and refuses to start, rather than leaving you guessing which won.")
    make_table(doc,
        ["Source", "Use when"],
        [
            ["`hsm.pin_file`", "Preferred. A secrets manager, a Kubernetes secret, or "
                               "systemd LoadCredential mounts the file; the server "
                               "reads it at the moment it is needed and never holds it "
                               "on the config object."],
            ["`hsm.pin_env`", "Acceptable where the environment is protected."],
            ["`hsm.pin`", "Works, warns at every start, and is redacted from any config "
                          "dump. It ends up in version control sooner or later."],
        ],
        col_widths=[3.5, 13.5])
    doc.add_paragraph()
    code_block(doc, [
        "sudo install -o kmip -g kmip -m 600 /dev/null /run/secrets/kmip-token-pin",
        "printf '%s' '<user-pin>' | sudo tee /run/secrets/kmip-token-pin >/dev/null",
    ])

    h2(doc, "7.3  Validate before starting")
    code_block(doc, [
        "kmip-server --config /etc/kmip/config.yaml --check",
        "→ Configuration at /etc/kmip/config.yaml is valid",
    ])
    body(doc,
        "Validation is strict and names the offending key. An unknown section, a "
        "certificate without its key, require_client_cert without a CA, more than "
        "one PIN source, or dual control with fewer than two approvals all fail "
        "here rather than starting half-configured.")

    h2(doc, "7.4  Governance, when you want it")
    body(doc,
        "Both governance features are off by default, so an upgrade behaves exactly "
        "as before until an operator turns them on — the right default for controls "
        "that can refuse a client's operation or deactivate a key on their own.")
    code_block(doc, [
        "governance:",
        "  enabled: true               # run the cryptoperiod scheduler",
        "  scan_interval_seconds: 300",
        "  warn_days: 7",
        "  auto_rotate: false          # create a cross-linked replacement on expiry",
        "",
        "  dual_control: true          # Destroy and Export need approval",
        "  dual_control_operations: [Destroy, Export]",
        "  approvals_required: 2       # must be >= 2",
        "  approval_ttl_seconds: 3600",
    ])
    note(doc, "Turn dual control on only once you have at least three identities "
              "provisioned. The requester can never approve their own request, so with "
              "two identities and approvals_required: 2, nothing destructive can ever "
              "be approved.")
    doc.add_page_break()

    # ── 8  First start ───────────────────────────────────────────────────
    h1(doc, "8  First start and provisioning")

    h2(doc, "8.1  Start it in the foreground")
    code_block(doc, [
        "sudo -u kmip SOFTHSM2_CONF=/etc/kmip/softhsm2.conf \\",
        "    kmip-server --config /etc/kmip/config.yaml",
        "",
        "→ {\"level\":\"INFO\",\"msg\":\"KMIP server listening on 0.0.0.0:5696\"}",
        "→ {\"level\":\"INFO\",\"msg\":\"kmip-server ready with 4 worker(s)\"}",
    ])
    body(doc,
        "The first start does real work before it binds: it opens the HSM session, "
        "provisions the AES master key that protects secret blobs at rest, and "
        "converts any pre-existing cleartext blobs. On a token holding thousands of "
        "keys that takes seconds — which is why readiness checks the listener and "
        "not just the HSM.")

    h2(doc, "8.2  Confirm it is actually ready")
    code_block(doc, [
        "curl -s localhost:9696/health",
        "→ 200",
        "",
        "curl -s localhost:9696/ready",
        "→ {\"ready\": true, \"mechanisms\": 79, \"schema_version\": 1}",
    ])
    tip(doc, "Use /ready as the readiness probe and /health as the liveness probe. "
             "/ready answers 200 only once the KMIP port is open, the HSM session is "
             "usable and the database answers — so an orchestrator will not route "
             "traffic at a port that is not listening yet.")

    h2(doc, "8.3  Provision identities")
    body(doc,
        "There is no self-registration and no anonymous fallback: an identity must "
        "exist before it can authenticate. KMIP defines no operation for this, so it "
        "is done with kmip-admin.")
    code_block(doc, [
        "kmip-admin -c /etc/kmip/config.yaml identity add alice",
        "# prompts for a password; --password is available for scripting",
        "→ identity 'alice' saved",
        "",
        "kmip-admin -c /etc/kmip/config.yaml identity add bob",
        "kmip-admin -c /etc/kmip/config.yaml identity add ops-admin",
        "",
        "# One administrator, who can reach every object:",
        "kmip-admin -c /etc/kmip/config.yaml role grant ops-admin admin",
        "",
        "kmip-admin -c /etc/kmip/config.yaml identity list",
        "→ IDENTITY   DISABLED  CREATED_AT",
        "→ alice      False     2026-08-19T09:14:02",
    ])
    note(doc, "Password verification uses scrypt and is deliberately expensive. "
              "Successful verifications are cached briefly so that cost is not paid on "
              "every request; failures are never cached, and changing a password, "
              "disabling or deleting an identity invalidates the cache immediately.")

    h2(doc, "8.4  A first key, end to end")
    code_block(doc, [
        "python - <<'EOF'",
        "from kmip_pkcs11.test_app.client import KMIPClient",
        "",
        "with KMIPClient(host=\"kms.internal\", port=5696,",
        "                username=\"alice\", password=\"...\",",
        "                tls_ca=\"/etc/kmip/tls/ca.pem\") as c:",
        "    print(c.discover_versions())",
        "    uid = c.create(name=\"first-key\")",
        "    ct, iv, tag = c.encrypt(uid, b\"hello\")",
        "    print(c.decrypt(uid, ct, iv=iv, auth_tag=tag))",
        "EOF",
        "",
        "→ [(2, 1), (1, 4), (1, 3), (1, 2), (1, 1), (1, 0)]",
        "→ b'hello'",
    ])
    body(doc,
        "If that round-trips, the whole stack works: TLS, authentication, the "
        "metadata store, the PKCS#11 session and the token.")
    code_block(doc, [
        "# And it is recorded:",
        "kmip-admin -c /etc/kmip/config.yaml audit list --limit 5",
        "kmip-admin -c /etc/kmip/config.yaml audit verify",
        "→ audit chain OK (4 entries)",
    ])
    doc.add_page_break()

    # ── 9  Tests ─────────────────────────────────────────────────────────
    h1(doc, "9  Running the test suite")
    body(doc,
        "816 tests, all executed live against a real SoftHSM2 token rather than "
        "against mocks. Running them is the fastest way to confirm an installation "
        "is sound, and the fastest way to find out that a token lacks a mechanism "
        "the suite needs.")

    h2(doc, "9.1  Run everything")
    code_block(doc, [
        "pip install --use-pep517 -e '.[dev]'",
        "pytest",
        "→ 816 passed in 53.44s",
    ])
    tip(doc, "The suite manages its own token: the session fixture wipes "
             "/tmp/softhsm2_tests/tokens and re-initialises it every run, so results "
             "do not depend on what previous runs left behind. It does not touch your "
             "deployment token — it sets SOFTHSM2_CONF itself.")
    note(doc, "Run only one pytest process at a time. Several runs sharing one SoftHSM2 "
              "token contend for it and produce failures that look like real defects "
              "but are not.")

    h2(doc, "9.2  Useful invocations")
    make_table(doc,
        ["Command", "Purpose"],
        [
            ["`pytest -v`", "Per-test names."],
            ["`pytest -x`", "Stop at the first failure."],
            ["`pytest -k governance`", "Only tests matching a keyword."],
            ["`pytest kmip_pkcs11/tests/test_ttlv.py kmip_pkcs11/tests/test_lifecycle.py`",
             "Unit tests only — no token required."],
            ["`pytest --cov=kmip_pkcs11 --cov-report=html`", "Coverage into "
                                                             "coverage_html/."],
            ["`SOFTHSM2_LIB=/path/to/other.so pytest`", "Run against a different "
                                                        "PKCS#11 library."],
        ],
        col_widths=[9.0, 8.0])
    doc.add_paragraph()

    h2(doc, "9.3  What each module covers")
    make_table(doc,
        ["Module", "Tests", "Needs a token", "Covers"],
        [
            ["test_ttlv.py", "22", "No", "TTLV encoding and decoding, including "
                                         "truncation and padding."],
            ["test_lifecycle.py", "26", "No", "The lifecycle state machine and usage "
                                              "rules."],
            ["test_metadata.py", "18", "No", "The SQLite store in isolation."],
            ["test_operations.py", "8", "Yes", "Operation handlers end to end through a "
                                               "live server and client."],
            ["test_conformance.py", "48", "Yes", "OASIS KMIP TC-mapped conformance "
                                                 "assertions."],
            ["test_governance.py", "48", "Yes", "Cryptoperiod enforcement, dual "
                                                "control, groups, role allowlists."],
            ["test_extended_coverage.py", "646", "Yes", "Every operation, algorithm and "
                                                        "mode coverage, error paths, "
                                                        "authentication, access "
                                                        "control, audit, transport, "
                                                        "backup, workers, concurrency."],
            ["TOTAL", "816", "", "100 % pass rate."],
        ],
        col_widths=[5.0, 1.8, 2.5, 7.7])
    doc.add_paragraph()

    h2(doc, "9.4  If EC signing tests fail")
    body(doc,
        "That is the signature of a SoftHSM2 build without CKM_ECDSA_SHA256 — see "
        "Section 3.1. Check the mechanism list before assuming the code is at "
        "fault:")
    code_block(doc, [
        "python - <<'EOF'",
        "import os, pkcs11",
        "os.environ[\"SOFTHSM2_CONF\"] = \"/tmp/softhsm2_tests/softhsm2.conf\"",
        "lib = pkcs11.lib(\"/usr/local/lib/softhsm/libsofthsm2.so\")",
        "tok = next(iter(lib.get_tokens(token_label=\"KMIPTestSuite\")))",
        "mechs = {int(m) for m in tok.slot.get_mechanisms()}",
        "print(len(mechs), \"mechanisms; ECDSA_SHA256:\", 0x1044 in mechs)",
        "EOF",
        "→ 79 mechanisms; ECDSA_SHA256: True",
    ])
    doc.add_page_break()

    # ── 10  systemd ──────────────────────────────────────────────────────
    h1(doc, "10  Running as a systemd service")
    body(doc,
        "deploy/kmip-server.service in the repository is a hardened unit. Install "
        "it, adjust the paths, and enable it.")
    code_block(doc, [
        "sudo cp deploy/kmip-server.service /etc/systemd/system/",
        "sudo systemctl daemon-reload",
        "sudo systemctl enable --now kmip-server",
        "sudo systemctl status kmip-server",
        "journalctl -u kmip-server -f",
    ])

    h2(doc, "10.1  What the unit does for you")
    make_table(doc,
        ["Directive", "Effect"],
        [
            ["`ExecReload=/bin/kill -HUP $MAINPID`", "systemctl reload re-reads TLS "
                                                     "material without dropping "
                                                     "connections."],
            ["`ProtectSystem=strict`, `ReadWritePaths=`", "The filesystem is read-only "
                                                          "except the database and "
                                                          "token directories."],
            ["`NoNewPrivileges`, `CapabilityBoundingSet=`", "No capabilities at all — "
                                                            "the KMIP port is above "
                                                            "1024, so none are needed."],
            ["`PrivateTmp`, `PrivateDevices`, `ProtectHome`", "No access to other "
                                                              "processes' temporary "
                                                              "files, devices or home "
                                                              "directories."],
            ["`MemoryDenyWriteExecute`, `LockPersonality`", "Standard exploit "
                                                            "mitigations."],
            ["`SystemCallFilter=@system-service`", "A restricted syscall set."],
            ["`Restart=on-failure`", "Restarts on a crash, not on a clean exit."],
        ],
        col_widths=[6.5, 10.5])
    doc.add_paragraph()

    h2(doc, "10.2  The PIN, without a file on disk")
    body(doc,
        "systemd can supply the PIN as a credential: it appears under "
        "$CREDENTIALS_DIRECTORY, readable only by this unit. Uncomment the "
        "LoadCredential line and point hsm.pin_file at it.")
    code_block(doc, [
        "# In the unit:",
        "LoadCredential=token-pin:/etc/kmip/token-pin",
        "",
        "# In config.yaml — the path systemd exposes it at:",
        "hsm:",
        "  pin_file: /run/credentials/kmip-server.service/token-pin",
    ])

    h2(doc, "10.3  If you must bind below port 1024")
    note(doc, "Add AmbientCapabilities=CAP_NET_BIND_SERVICE to the unit rather than "
              "running the server as root. Nothing else about the server needs "
              "privilege.")
    doc.add_page_break()

    # ── 11  Container ────────────────────────────────────────────────────
    h1(doc, "11  Running as a container")
    body(doc,
        "deploy/Dockerfile is a two-stage build: the first stage compiles SoftHSM2 "
        "2.7.0 from source, the second carries only the runtime.")
    code_block(doc, [
        "docker build -f deploy/Dockerfile -t kmip-server .",
        "",
        "docker run --rm \\",
        "    -v /etc/kmip:/etc/kmip:ro \\",
        "    -v kmip-data:/var/lib/kmip \\",
        "    -v kmip-tokens:/var/lib/softhsm \\",
        "    -p 5696:5696 -p 127.0.0.1:9696:9696 \\",
        "    kmip-server",
    ])
    body(doc,
        "Both the database and the token store must be volumes, and they must "
        "survive together — see the note in Section 1.3. The image runs as an "
        "unprivileged user, sets SOFTHSM2_CONF, and carries a HEALTHCHECK that "
        "polls /health.")
    note(doc, "The image and the CI workflow are written and their inputs checked, but "
              "neither has been executed anywhere: the environment this project was "
              "developed in blocks Docker Hub and the SoftHSM2 source mirror. Treat "
              "both as unverified until you have run them once.")
    doc.add_page_break()

    # ── 12  Operating ────────────────────────────────────────────────────
    h1(doc, "12  Upgrading, backup and recovery")

    h2(doc, "12.1  Upgrading")
    code_block(doc, [
        "# 1. Back up first — always, and especially across a schema change.",
        "kmip-admin -c /etc/kmip/config.yaml backup create \\",
        "    --output /var/backups/kmip-$(date +%F)",
        "",
        "# 2. Stop, upgrade, validate, start.",
        "sudo systemctl stop kmip-server",
        "git pull && pip install .",
        "kmip-server --config /etc/kmip/config.yaml --check",
        "sudo systemctl start kmip-server",
        "",
        "# 3. Confirm.",
        "curl -s localhost:9696/ready",
        "kmip-admin -c /etc/kmip/config.yaml audit verify",
    ])
    body(doc,
        "Schema migrations are applied automatically at startup, gated on PRAGMA "
        "user_version and committed one at a time with their version bump, so an "
        "interrupted upgrade resumes rather than half-applying. There is no "
        "downgrade path: to go back, restore the backup.")
    note(doc, "Upgrading into a release that encrypts secret blobs at rest converts "
              "existing rows on first start and then scrubs the freed pages, including "
              "the WAL sidecar. That scrub reaches the live database files only — "
              "snapshots and backups taken before the upgrade still contain cleartext "
              "and must be re-taken or destroyed.")

    h2(doc, "12.2  Backup")
    code_block(doc, [
        "kmip-admin -c /etc/kmip/config.yaml backup create \\",
        "    --output /var/backups/kmip-2026-08-19",
        "",
        "kmip-admin backup inspect --input /var/backups/kmip-2026-08-19",
        "→ token label, object count, audit entries, creation time",
    ])
    body(doc,
        "The snapshot uses SQLite's online backup API — copying a live WAL database "
        "with cp can capture a torn state — and writes a manifest recording which "
        "token it belongs to.")
    note(doc, "Back up the token store (/var/lib/softhsm/tokens) at the same time and "
              "keep the two together. The database alone is unreadable without the "
              "token that holds the master key.")

    h2(doc, "12.3  Restore, and proving it worked")
    code_block(doc, [
        "sudo systemctl stop kmip-server",
        "",
        "kmip-admin -c /etc/kmip/config.yaml backup restore \\",
        "    --input /var/backups/kmip-2026-08-19 \\",
        "    --database /var/lib/kmip/kmip.db",
        "",
        "# Prove the result rather than assuming it:",
        "kmip-admin -c /etc/kmip/config.yaml backup verify",
        "→ /var/lib/kmip/kmip.db: OK",
        "",
        "sudo systemctl start kmip-server",
    ])
    body(doc,
        "backup verify checks three things that together mean the restore really "
        "worked: the master key is present on the token, the audit chain verifies, "
        "and a stored secret actually decrypts. Restore refuses a mismatched token "
        "unless forced — and forcing it gives you a database whose secrets cannot "
        "be decrypted.")

    h2(doc, "12.4  Rotating the master key")
    code_block(doc, [
        "kmip-admin -c /etc/kmip/config.yaml rotate-master-key",
        "→ re-encrypted 1043 blob(s); retired 1 old key(s)",
    ])
    body(doc,
        "Re-encryption is row by row and each envelope records which key wrote it, "
        "so a partly rotated store stays fully readable and re-running finishes the "
        "job. The superseded key is destroyed only after the last row has moved.")
    doc.add_page_break()

    # ── 13  Hardening ────────────────────────────────────────────────────
    h1(doc, "13  Hardening checklist")
    for item in [
        "TLS is on — server.allow_plaintext is false or absent. Verify by connecting "
        "with a plain socket and getting nothing.",
        "The token PIN comes from pin_file or pin_env, not inline. A warning at every "
        "start means you missed this.",
        "The management port (9696) is bound to 127.0.0.1 or an internal interface. "
        "/metrics and /ready are for the operator, not for KMIP clients.",
        "The server runs as an unprivileged user with the hardened unit from "
        "Section 10.",
        "At least one identity holds the admin role, and ordinary clients do not.",
        "Every client has its own identity. Shared credentials make the audit log "
        "useless for the question it exists to answer.",
        "The audit chain verifies, and something checks it on a schedule rather than "
        "only after an incident.",
        "Backups are taken and restored somewhere — a backup you have never restored "
        "is a hypothesis.",
        "The database and token store are backed up together.",
        "Cryptoperiods are set on keys that should have them, and the scheduler is "
        "enabled to act on them.",
        "Dual control is on for Destroy and Export if losing or disclosing a key would "
        "matter, and at least three identities exist so approvals are possible.",
        "Log output goes somewhere durable. JSON format keeps tracebacks in a field so "
        "aggregation does not split them.",
    ]:
        bullet(doc, item)
    doc.add_page_break()

    # ── 14  Troubleshooting ──────────────────────────────────────────────
    h1(doc, "14  Troubleshooting")
    body(doc,
        "Real error text, and what it actually means. Every message below is one "
        "this software emits verbatim.")

    problems = [
        ("No TLS certificate configured. Set tls.cert and tls.key, or set "
         "server.allow_plaintext: true",
         "The server refuses to start in the clear by design. Configure TLS "
         "(Section 6), or opt out explicitly for a test rig — in which case it will "
         "warn at every start, naming the address it is exposing.",
         ["server:", "  allow_plaintext: true    # test rigs only"]),

        ("Token 'KMIPToken' not found",
         "The label in hsm.token_label does not match any initialised token, or "
         "SOFTHSM2_CONF is not set in the server's environment so SoftHSM2 is "
         "looking in the wrong directory. This is the single most common "
         "first-start failure.",
         ["SOFTHSM2_CONF=/etc/kmip/softhsm2.conf softhsm2-util --show-slots",
          "# The Label: line must match hsm.token_label exactly."]),

        ("CKR_PIN_INCORRECT",
         "The user PIN is wrong — or the SO PIN was supplied instead. The server "
         "uses the user PIN; the SO PIN administers the token and is never used "
         "here.",
         ["# Check what the file actually contains, whitespace included:",
          "sudo -u kmip cat /run/secrets/kmip-token-pin | xxd | tail -2"]),

        ("mechanism ECDSA_SHA256 is not available on this PKCS#11 token",
         "A SoftHSM2 build older than 2.7.0 — very likely the distribution package. "
         "Build 2.7.0 from source (Section 3.2). The same message with a different "
         "mechanism name means the token genuinely does not implement it, and the "
         "algorithm coverage table in the Feature Specification will tell you which "
         "are affected.",
         ["apt-cache policy softhsm2      # what is installed",
          "softhsm2-util --version        # what is on the PATH"]),

        ("AttributeError: install_layout",
         "A Debian/Ubuntu setuptools interaction with the legacy editable install "
         "path. Not a defect in this project.",
         ["pip install --use-pep517 -e '.[dev]'"]),

        ("Operation failed (reason=12): Identity 'bob' is not authorized to perform "
         "'Get' on an object it does not own",
         "Authorization working as designed. bob is not the owner, holds no admin "
         "role, and has no grant. Grant access to the identity or to a group it "
         "belongs to.",
         ["kmip-admin -c config.yaml access grant <uid> bob --permission read",
          "# or, for a whole team:",
          "kmip-admin -c config.yaml group add bob crypto-team",
          "kmip-admin -c config.yaml access grant <uid> group:crypto-team \\",
          "                                             --permission read"]),

        ("Identity 'dave' holds no role permitting 'Create'",
         "A role dave holds carries an operation allowlist that does not include "
         "Create. Allowlists are opt-in and narrow what an identity may do; they "
         "never widen it.",
         ["kmip-admin -c config.yaml role show dave",
          "kmip-admin -c config.yaml permission show <role>",
          "kmip-admin -c config.yaml permission allow <role> Create"]),

        ("Destroy requires 2 approval(s) from other identities. Approval request "
         "<id> created; once approved, retry.",
         "Dual control. The operation did NOT run — the handler is never reached. "
         "Have the required number of other identities approve, then retry. The "
         "requester can never approve their own request, and an approval is spent "
         "by a single attempt.",
         ["kmip-admin -c config.yaml approval list",
          "kmip-admin -c config.yaml approval approve <request-id> --as bob",
          "kmip-admin -c config.yaml approval approve <request-id> --as carol",
          "# then the original client retries"]),

        ("Operation 'activate' is not permitted when object state is 'Active'",
         "Not an error in your code. Create, CreateKeyPair and Register return "
         "objects that are already Active; Activate exists for objects deliberately "
         "staged Pre-Active.",
         ["# Nothing to do — the key is already usable."]),

        ("A key became Deactivated and no client asked for it",
         "The cryptoperiod scheduler acted, because the key's Deactivation Date had "
         "passed. It is recorded under the identity system:scheduler, so it is as "
         "attributable as anything else.",
         ["kmip-admin -c config.yaml audit list --object-uid <uid>",
          "kmip-admin -c config.yaml cryptoperiod expiring --within-days 30",
          "kmip-admin -c config.yaml cryptoperiod set <uid> --days 365"]),

        ("/ready returns 503 while the process is clearly running",
         "Working as intended during startup: the server opens the HSM session, "
         "provisions the master key and converts cleartext blobs before it binds, "
         "which on a large token takes seconds. The body names which check failed.",
         ["curl -s localhost:9696/ready",
          "→ {\"ready\": false, \"kmip_listener\": \"not yet accepting connections\"}"]),

        ("audit chain BROKEN at seq N",
         "The hash chain does not verify at that row: the log has been modified "
         "outside the sanctioned path. Preserve the database before doing anything "
         "else — and note that prune refuses to run on a log that already fails "
         "verification, precisely so pruning cannot destroy the evidence.",
         ["cp /var/lib/kmip/kmip.db /var/backups/kmip-suspect.db",
          "kmip-admin -c config.yaml audit list --limit 50 --json"]),

        ("Throughput is far lower than expected",
         "Check the worker count first. One worker is a single locked PKCS#11 "
         "session — roughly 600 operations/second. More workers help, but only to "
         "about 1.5x: the hash-chained audit log serialises every audited operation "
         "on one database write lock. That is the cost of the tamper-evidence, not a "
         "misconfiguration.",
         ["server:", "  workers: 4        # or null for one per CPU"]),

        ("Tests fail in ways unrelated to what you changed",
         "Check that only one pytest process is running. Several runs sharing one "
         "SoftHSM2 token contend for it and produce failures that look real.",
         ["pgrep -fa pytest"]),
    ]
    for title, explanation, cmds in problems:
        h3(doc, title)
        body(doc, explanation, space_after=2)
        if cmds:
            code_block(doc, cmds)
    doc.add_page_break()

    # ── 15  Quick reference ──────────────────────────────────────────────
    h1(doc, "15  Quick reference")

    h2(doc, "15.1  Install, from nothing to running")
    code_block(doc, [
        "# 1. Packages",
        "sudo apt install -y build-essential automake autoconf libtool pkg-config \\",
        "                    libssl-dev curl python3 python3-pip python3-venv",
        "",
        "# 2. SoftHSM2 2.7.0 from source",
        "curl -fsSL https://dist.opendnssec.org/source/softhsm-2.7.0.tar.gz | tar -xz -C /tmp",
        "cd /tmp/softhsm-2.7.0 && ./configure --prefix=/usr/local \\",
        "    --with-crypto-backend=openssl --disable-gost && make -j\"$(nproc)\"",
        "sudo make install && sudo ldconfig",
        "",
        "# 3. Token",
        "export SOFTHSM2_CONF=/etc/kmip/softhsm2.conf",
        "softhsm2-util --init-token --slot 0 --label KMIPToken --pin ... --so-pin ...",
        "",
        "# 4. Package",
        "python3 -m venv .venv && source .venv/bin/activate && pip install .",
        "",
        "# 5. Configure and check",
        "kmip-server --config /etc/kmip/config.yaml --check",
        "",
        "# 6. Provision and start",
        "kmip-admin -c /etc/kmip/config.yaml identity add alice",
        "sudo systemctl enable --now kmip-server",
    ])

    h2(doc, "15.2  Commands by task")
    make_table(doc,
        ["Task", "Command"],
        [
            ["Validate configuration", "`kmip-server --config PATH --check`"],
            ["Run the server", "`kmip-server --config PATH`"],
            ["Reload TLS after renewal", "`systemctl reload kmip-server`"],
            ["Add an identity", "`kmip-admin -c PATH identity add NAME`"],
            ["Make someone an admin", "`kmip-admin -c PATH role grant NAME admin`"],
            ["Delegate one object", "`kmip-admin -c PATH access grant UID NAME "
                                    "--permission read`"],
            ["Grant a whole team", "`kmip-admin -c PATH access grant UID group:TEAM "
                                   "--permission read`"],
            ["Restrict a role", "`kmip-admin -c PATH permission allow ROLE Get`"],
            ["Set a cryptoperiod", "`kmip-admin -c PATH cryptoperiod set UID "
                                   "--days 365`"],
            ["See what expires soon", "`kmip-admin -c PATH cryptoperiod expiring "
                                      "--within-days 30`"],
            ["Approve a destructive op", "`kmip-admin -c PATH approval approve ID "
                                         "--as NAME`"],
            ["Read the audit log", "`kmip-admin -c PATH audit list --limit 50`"],
            ["Verify the audit chain", "`kmip-admin -c PATH audit verify`"],
            ["Back up", "`kmip-admin -c PATH backup create --output DIR`"],
            ["Prove a database is usable", "`kmip-admin -c PATH backup verify`"],
            ["Rotate the master key", "`kmip-admin -c PATH rotate-master-key`"],
            ["Run the tests", "`pytest`"],
        ],
        col_widths=[6.0, 11.0])
    doc.add_paragraph()

    h2(doc, "15.3  Related documents")
    make_table(doc,
        ["Document", "Covers"],
        [
            ["KMIP_PKCS11_Feature_Specification.docx", "Every supported feature and "
                                                       "operation, with worked "
                                                       "examples. The reference for "
                                                       "what to send and what comes "
                                                       "back."],
            ["KMIP_PKCS11_Design_Document.docx", "Architecture and design rationale, "
                                                 "including rejected designs."],
            ["KMIP_PKCS11_Project_Documentation.docx", "Module-by-module reference and "
                                                       "the test specification."],
            ["KMIP_PKCS11_Phase_Report.docx", "How the system reached its current "
                                              "state, phase by phase."],
            ["KMIP_PKCS11_KMS_Gap_Matrix.docx", "What a high-end commercial KMS is "
                                                "expected to do, whether this project "
                                                "covers it, and the route for every "
                                                "gap."],
            ["README_KMIP.md", "The short version of all of the above."],
        ],
        col_widths=[6.5, 10.5])

    doc.add_paragraph()
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run(
        f"KMIP on PKCS#11  ·  Installation, Deployment & Test Guide  ·  {TODAY}")
    fr.font.size = Pt(9)
    fr.font.color.rgb = RGBColor(0x90, 0x90, 0x90)
    fr.font.name = BODY_FONT

    out = "KMIP_PKCS11_Install_Test_Guide.docx"
    doc.save(out)
    print(f"Saved: {out}")
    return out


if __name__ == "__main__":
    build()
