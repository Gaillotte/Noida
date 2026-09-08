"""Render a .pptx to PNGs with PIL, so slides can be inspected where
LibreOffice cannot run.

Not a faithful PowerPoint renderer — it draws shape geometry, fills and
wrapped text at their real sizes and positions. That is enough to see what a
render would show: missing shapes, overlaps, overflow, and whether a diagram
reads as a diagram.
"""
import sys, os
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

DPI = 110
SW, SH = 13.333, 7.5
FONT_DIR = "/usr/share/fonts/truetype/dejavu"

def px(v): return int(round(v * DPI))

def font(size_pt, bold=False, italic=False):
    # Only Regular and Bold ship here; italic falls back to regular, which is
    # fine for judging layout.
    name = "DejaVuSans-Bold" if bold else "DejaVuSans"
    return ImageFont.truetype(os.path.join(FONT_DIR, name + ".ttf"),
                              max(6, int(round(size_pt * DPI / 72.0))))

def rgb(color_obj, default=None):
    try:
        if color_obj and color_obj.type is not None and color_obj.rgb is not None:
            return "#" + str(color_obj.rgb)
    except Exception:
        pass
    return default

def shape_fill(sh):
    try:
        if sh.fill.type is not None and sh.fill.type == 1:      # solid
            return rgb(sh.fill.fore_color)
    except Exception:
        pass
    return None

def shape_line(sh):
    try:
        return rgb(sh.line.color)
    except Exception:
        return None

def wrap(draw, text, fnt, max_w):
    out = []
    for para in text.split("\n"):
        if not para:
            out.append("")
            continue
        words, line = para.split(" "), ""
        for w in words:
            t = (line + " " + w).strip()
            if draw.textlength(t, font=fnt) <= max_w or not line:
                line = t
            else:
                out.append(line); line = w
        out.append(line)
    return out

def render(path, prefix):
    prs = Presentation(path)
    made = []
    for n, slide in enumerate(prs.slides, 1):
        bg = "#FFFFFF"
        try:
            if slide.background.fill.type == 1:
                bg = rgb(slide.background.fill.fore_color, "#FFFFFF")
        except Exception:
            pass
        img = Image.new("RGB", (px(SW), px(SH)), bg)
        d = ImageDraw.Draw(img)

        for sh in slide.shapes:
            if sh.left is None: continue
            x, y = Emu(sh.left).inches, Emu(sh.top).inches
            w, h = Emu(sh.width).inches, Emu(sh.height).inches
            box = [px(x), px(y), px(x + w), px(y + h)]

            if sh.shape_type is not None and "TABLE" in str(sh.shape_type):
                pass                                    # drawn below
            elif sh.has_text_frame and not sh.shape_type == 1:
                pass

            fill = shape_fill(sh)
            ln = shape_line(sh)
            is_txtbox = str(sh.shape_type).startswith("TEXT_BOX")
            if (fill or ln) and not is_txtbox:
                if box[2] - box[0] < 2: box[2] = box[0] + 2   # keep hairlines visible
                if box[3] - box[1] < 2: box[3] = box[1] + 2
                d.rounded_rectangle(box, radius=6, fill=fill, outline=ln or fill, width=1)

            if sh.has_table:
                t = sh.table
                rows = len(t.rows); cols = len(t.columns)
                colw = [Emu(c.width).inches for c in t.columns]
                rowh = [Emu(r.height).inches for r in t.rows]
                cy = y
                for ri in range(rows):
                    cx = x
                    for ci in range(cols):
                        cell = t.cell(ri, ci)
                        cb = [px(cx), px(cy), px(cx + colw[ci]), px(cy + rowh[ri])]
                        cf = None
                        try:
                            if cell.fill.type == 1: cf = rgb(cell.fill.fore_color)
                        except Exception: pass
                        d.rectangle(cb, fill=cf, outline="#D6DEE7")
                        runs = [r for p in cell.text_frame.paragraphs for r in p.runs]
                        sz = runs[0].font.size.pt if runs and runs[0].font.size else 11
                        col = rgb(runs[0].font.color, "#16202B") if runs else "#16202B"
                        bold = bool(runs and runs[0].font.bold)
                        f = font(sz, bold)
                        d.text((cb[0] + 5, cb[1] + 4), cell.text, font=f, fill=col)
                        cx += colw[ci]
                    cy += rowh[ri]
                continue

            if not sh.has_text_frame: continue
            tf = sh.text_frame
            if not tf.text.strip(): continue

            ty = px(y) + 3
            for p in tf.paragraphs:
                if not p.runs:
                    ty += px(0.12); continue
                r0 = p.runs[0]
                sz = r0.font.size.pt if r0.font.size else 12
                col = rgb(r0.font.color, "#16202B")
                f = font(sz, bool(r0.font.bold), bool(r0.font.italic))
                text = "".join(r.text for r in p.runs)
                bullet = "•  " if p.level == 0 and _has_bullet(p) else ""
                lines = wrap(d, bullet + text, f, px(w) - 8)
                for ln_txt in lines:
                    tw = d.textlength(ln_txt, font=f)
                    tx = px(x) + 4
                    align = str(p.alignment or "")
                    if "CENTER" in align: tx = px(x) + (px(w) - tw) / 2
                    elif "RIGHT" in align: tx = px(x + w) - tw - 4
                    d.text((tx, ty), ln_txt, font=f, fill=col)
                    ty += f.size * 1.22

        out = f"{prefix}-{n:02d}.png"
        img.save(out)
        made.append(os.path.abspath(out))
    return made

def _has_bullet(p):
    xml = p._p.xml
    return "buChar" in xml or "buAutoNum" in xml

if __name__ == "__main__":
    for f in render(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "render"):
        print(f)
