"""
Generate the provenance matrix: which source evidenced each requirement in the
gap analysis.

Run: python generate_feature_provenance.py

READ THIS BEFORE READING THE TABLE.

A mark says "this source's documentation evidenced this requirement". It does
NOT say "this product has this feature", and a blank does NOT say a product
lacks it. The research behind the gap matrix was a handful of published pages
per vendor — enough to establish what the market expects a KMS to do, nowhere
near enough to compare products feature by feature. A capability comparison
would need a hands-on evaluation of each product and is not what this is.

Marks are given strictly. A source that documents "audit logging" evidences an
audit trail, not a tamper-evident one, so it does not mark the tamper-evidence
row. Blank cells in rows a product plainly satisfies are the expected cost of
that rule.

CAT is the residue: requirements no source in the research named, present
because they are ordinary practice for a networked service handling secrets.
It is mutually exclusive with every other code — verify() enforces that — so
the CAT column count is exactly the number of unsourced requirements.
"""

import datetime
import importlib.util
import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DARK_BLUE = RGBColor(0x1A, 0x3A, 0x5C)
MID_BLUE = RGBColor(0x2E, 0x6D, 0xA4)
DARK_GREY = RGBColor(0x40, 0x40, 0x40)
OK_GREEN = RGBColor(0x1E, 0x80, 0x2E)
WARN_AMBER = RGBColor(0x8F, 0x56, 0x10)
GAP_RED = RGBColor(0xB3, 0x26, 0x1E)

TODAY = datetime.date.today().strftime("%d %B %Y")

# The sources, in the order they appear as columns.
SOURCES = [
    ("THA", "Thales CipherTrust Manager", "vendor"),
    ("FOR", "Fortanix Data Security Manager", "vendor"),
    ("ENT", "Entrust KeyControl", "vendor"),
    ("UTI", "Utimaco Enterprise Secure Key Manager", "vendor"),
    ("BLO", "Bloombase KeyCastle", "vendor"),
    ("COS", "Cosmian / Eviden KMS", "vendor"),
    ("SEC", "Securosys CyberVault KMS", "vendor"),
    ("VLT", "HashiCorp Vault", "vendor"),
    ("CLD", "AWS KMS / Azure Managed HSM / Google Cloud KMS", "vendor"),
    ("STD", "NIST SP 800-57, SP 800-152, OASIS KMIP 2.1 and 3.0", "standard"),
    ("CAT", "Category baseline — no source named it", "residue"),
]

# feature name (must match the gap matrix exactly) -> sources that evidenced it
PROVENANCE = {
    # ── Key lifecycle and cryptography ───────────────────────────────────
    "HSM-backed storage, keys never leave": ["THA", "FOR", "ENT", "UTI", "BLO",
                                             "COS", "SEC", "CLD", "STD"],
    "Full lifecycle state machine": ["THA", "ENT", "BLO", "COS", "SEC", "STD"],
    "Cryptoperiod enforcement": ["ENT", "CLD", "STD"],
    "Automatic rotation": ["THA", "ENT", "UTI", "COS", "SEC", "CLD"],
    "Key lineage and versioning": ["THA", "VLT", "CLD"],
    "Algorithm breadth": ["COS", "SEC", "VLT", "STD"],
    "Post-quantum (ML-KEM, ML-DSA, SLH-DSA)": ["COS", "SEC", "STD"],
    "Hardware RNG": ["UTI", "BLO", "STD"],
    "Key derivation": ["VLT", "STD"],
    "Key wrapping / secure import-export": ["THA", "BLO", "VLT", "CLD", "STD"],
    "Split knowledge / M-of-N shares": ["FOR", "UTI", "VLT", "STD"],
    "Formal key ceremony": ["VLT", "STD"],
    "Bulk and batch operations": ["STD"],
    "Verifiable destruction": ["THA", "ENT", "STD"],
    "Key-bound usage policy enforced in hardware": ["SEC"],

    # ── Protocol and ecosystem integration ───────────────────────────────
    "KMIP server": ["THA", "ENT", "UTI", "BLO", "COS", "SEC", "VLT", "STD"],
    "KMIP interoperability certification": ["BLO", "STD"],
    "KMIP 3.0": ["STD"],
    "PKCS#11 provider for applications": ["UTI", "BLO", "COS", "SEC"],
    "REST / JSON API": ["THA", "FOR", "UTI", "COS", "SEC", "VLT", "CLD"],
    "Cloud BYOK (AWS, Azure, GCP)": ["THA", "FOR", "CLD"],
    "Cloud EKM / XKS / hold-your-own-key": ["FOR", "SEC", "CLD"],
    "Database TDE integration": ["THA", "ENT", "COS", "SEC"],
    "Storage, backup and VM integrations": ["ENT", "BLO", "SEC"],
    "Secrets management": ["THA", "FOR", "ENT", "VLT"],
    "Tokenization / format-preserving encryption": ["THA", "FOR", "COS"],
    "Certificate lifecycle": ["BLO", "COS", "SEC"],
    "Encryption as a service (data-plane API)": ["COS", "VLT", "CLD"],
    "KMIP JSON and XML encodings": ["COS", "STD"],
    "OpenAPI specification and generated clients": ["COS", "CLD"],

    # ── Identity and access control ──────────────────────────────────────
    "Per-identity authentication": ["THA", "UTI", "COS", "VLT"],
    "Enterprise IdP — LDAP/AD, SAML, OIDC": ["UTI", "COS", "VLT"],
    "mTLS certificate identity": ["COS", "VLT"],
    "API keys / service accounts": ["COS", "VLT", "CLD"],
    "Role-based access control": ["THA", "UTI", "SEC", "VLT"],
    "Groups": ["VLT"],
    "Per-object access grants": ["THA", "COS"],
    "Attribute or policy-based access": ["THA", "FOR", "ENT", "COS", "SEC"],
    "Dual control / quorum approval": ["FOR", "UTI", "SEC", "VLT"],
    "Separation of duties": ["STD"],
    "Multi-tenancy and namespaces": ["THA", "ENT", "SEC", "VLT"],
    "Per-tenant quotas": ["VLT"],
    "Connection and rate limiting": ["VLT"],

    # ── Audit, compliance and assurance ──────────────────────────────────
    "Tamper-evident audit log": ["THA", "UTI", "STD"],
    "External audit anchoring": ["CAT"],
    "SIEM integration": ["UTI", "BLO", "VLT"],
    "Retention and archival": ["BLO", "STD"],
    "FIPS 140-2/3 validated HSM": ["FOR", "ENT", "BLO", "COS", "SEC", "CLD"],
    "Common Criteria / eIDAS": ["ENT", "SEC"],
    "Compliance reporting": ["THA", "ENT", "SEC"],
    "Key inventory and discovery": ["THA", "ENT", "SEC"],
    "Cryptoperiod policy alignment": ["STD"],
    "Crypto-agility reporting": ["SEC"],
    "Certificate discovery and expiry monitoring": ["SEC"],

    # ── Availability, scale and recovery ─────────────────────────────────
    "HA clustering": ["THA", "UTI", "BLO", "COS", "VLT"],
    "Multi-site replication and DR": ["ENT", "UTI", "VLT", "STD"],
    "HSM failover and pooling": ["UTI", "BLO"],
    "Horizontal throughput": ["THA", "COS", "VLT"],
    "Enterprise database backend": ["COS", "VLT"],
    "Scheduled and offsite backup": ["THA", "ENT", "STD"],
    "Restore that proves itself": ["STD"],
    "Zero-downtime certificate rotation": ["CAT"],
    "Rolling upgrades": ["CAT"],
    "Kubernetes / container deployment": ["THA", "COS", "SEC"],
    "Confidential computing deployment": ["COS"],

    # ── Operations and observability ─────────────────────────────────────
    "Health and readiness probes": ["CAT"],
    "Metrics": ["VLT"],
    "Structured logging": ["CAT"],
    "Configuration-driven deployment": ["CAT"],
    "Administrative CLI": ["COS", "VLT"],
    "Web console": ["THA", "FOR", "ENT", "BLO", "COS", "SEC"],
    "Self-service developer portal": ["THA"],
    "Alerting": ["BLO", "SEC"],
    "Distributed tracing": ["COS"],

    # ── Platform hardening ───────────────────────────────────────────────
    "TLS enforced by default": ["CAT"],
    "Metadata encrypted at rest": ["COS", "VLT", "STD"],
    "Master key rotation": ["VLT", "STD"],
    "Non-extractable key enforcement": ["SEC", "CLD", "STD"],
    "Request bounds and DoS resistance": ["CAT"],
    "Secrets kept out of configuration": ["CAT"],
    "Hardened service": ["CAT"],
    "Sealed startup with quorum unseal": ["VLT"],
}

TOTAL = 0          # set from the matrix at build time; never hard-coded

STATUS_LABEL = {"full": "Covered", "partial": "Partial", "gap": "Gap"}
STATUS_COLOUR = {"full": OK_GREEN, "partial": WARN_AMBER, "gap": GAP_RED}


def load_matrix():
    here = os.path.dirname(os.path.abspath(__file__))
    spec = importlib.util.spec_from_file_location(
        "gapmatrix", os.path.join(here, "generate_kms_gap_matrix.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def verify(matrix):
    """Same discipline as the delivery plan: this table must describe the
    matrix it claims to, or it does not get built."""
    listed = [r[0] for d in matrix.DOMAINS for r in d["rows"]]
    keys = set(PROVENANCE)

    unknown = sorted(keys - set(listed))
    if unknown:
        raise SystemExit("provenance names features the matrix does not have:\n  "
                         + "\n  ".join(unknown))
    missing = sorted(set(listed) - keys)
    if missing:
        raise SystemExit("features with no provenance entry:\n  " + "\n  ".join(missing))

    codes = {c for c, _n, _k in SOURCES}
    for feat, marks in PROVENANCE.items():
        bad = [m for m in marks if m not in codes]
        if bad:
            raise SystemExit(f"{feat}: unknown source code(s) {bad}")
        if not marks:
            raise SystemExit(f"{feat}: no source at all — every requirement needs one")
        if len(set(marks)) != len(marks):
            raise SystemExit(f"{feat}: duplicate source code")
        if "CAT" in marks and len(marks) > 1:
            raise SystemExit(
                f"{feat}: CAT means no source named it, so it cannot sit beside "
                f"{[m for m in marks if m != 'CAT']}")
    return listed


# ── document helpers ─────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_heading(doc, text, level, colour=None):
    h = doc.add_heading(text, level=level)
    if colour:
        for run in h.runs:
            run.font.color.rgb = colour
    return h


def add_para(doc, text="", bold=False, italic=False, size=10, colour=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if colour:
        run.font.color.rgb = colour
    return p


def add_bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.32)
    p.add_run(text).font.size = Pt(size)
    return p


def section_break(doc):
    doc.add_paragraph()


COLS = [c for c, _n, _k in SOURCES]
# A4 landscape leaves 10.59" between the margins. Spend it on the requirement
# name first — a wrapped feature name is much harder to read than a tight tick
# column — then divide the rest.
_SRC_W = 0.50
COL_W = [10.59 - _SRC_W * len(COLS) - 0.80] + [_SRC_W] * len(COLS) + [0.80]


def matrix_header(tbl):
    hdr = tbl.rows[0]
    labels = ["Requirement"] + COLS + ["Status"]
    for i, text in enumerate(labels):
        cell = hdr.cells[i]
        cell.text = text
        set_cell_bg(cell, "1A3A5C")
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(8)


def matrix_row(tbl, feature, marks, status, shade):
    row = tbl.add_row()
    c0 = row.cells[0]
    r0 = c0.paragraphs[0].add_run(feature)
    r0.font.size = Pt(8)
    set_cell_bg(c0, shade)

    for i, code in enumerate(COLS, start=1):
        cell = row.cells[i]
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hit = code in marks
        run = para.add_run("●" if hit else "·")
        run.font.size = Pt(10 if hit else 8)
        # The category column is the residue, not a source; colour it as such
        # so a row carried only by it reads as weaker at a glance.
        run.font.color.rgb = (WARN_AMBER if (hit and code == "CAT")
                              else MID_BLUE if hit
                              else RGBColor(0xC8, 0xD0, 0xD8))
        set_cell_bg(cell, shade)

    cs = row.cells[len(COLS) + 1]
    p = cs.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(STATUS_LABEL[status])
    r.font.size = Pt(7.5)
    r.bold = True
    r.font.color.rgb = STATUS_COLOUR[status]
    set_cell_bg(cs, shade)


def fixed_layout(tbl):
    """Without this the renderer autofits and gives eleven one-character tick
    columns more width than the requirement names, which is how the first
    render of this table came out."""
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = tbl._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, w in zip(grid.findall(qn("w:gridCol")), COL_W):
            col.set(qn("w:w"), str(int(w * 1440)))


def widths(tbl):
    fixed_layout(tbl)
    for row in tbl.rows:
        for cell, w in zip(row.cells, COL_W):
            cell.width = Inches(w)


def build():
    matrix = load_matrix()
    listed = verify(matrix)
    global TOTAL
    TOTAL = len(listed)

    doc = Document()
    for sec in doc.sections:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)   # A4 landscape
        sec.top_margin = Inches(0.55)
        sec.bottom_margin = Inches(0.5)
        sec.left_margin = Inches(0.55)
        sec.right_margin = Inches(0.55)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    # ── header ───────────────────────────────────────────────────────────
    t = doc.add_paragraph()
    tr = t.add_run(f"Requirement provenance — where each of the {TOTAL} "
                   f"features came from")
    tr.bold = True
    tr.font.size = Pt(19)
    tr.font.color.rgb = DARK_BLUE

    add_para(doc,
        "A mark means the source's published documentation evidenced this "
        "requirement. It does NOT mean the product has the feature, and a blank "
        "does NOT mean it lacks one.", bold=True, size=10, colour=GAP_RED)
    add_para(doc,
        "The research behind the gap matrix was a handful of published pages per "
        "vendor — enough to establish what the market expects a key management "
        "system to do, nowhere near enough to compare products feature by "
        "feature. A capability comparison would need a hands-on evaluation of "
        "each product, and this is not one.", size=9.5)
    section_break(doc)

    # legend
    # Eleven codes will not read across a single row, so the legend runs down
    # the page in two columns of pairs.
    pairs = [SOURCES[i:i + 2] for i in range(0, len(SOURCES), 2)]
    legend = doc.add_table(rows=len(pairs), cols=4)
    legend.style = "Table Grid"
    for r, pair in enumerate(pairs):
        for c, entry in enumerate(pair):
            code, name, kind = entry
            shade = "FBF3E3" if kind == "residue" else "EEF4FA"
            cc = legend.rows[r].cells[c * 2]
            cn = legend.rows[r].cells[c * 2 + 1]
            cc.width, cn.width = Inches(0.62), Inches(4.67)
            rc = cc.paragraphs[0].add_run(code)
            rc.bold = True
            rc.font.size = Pt(9)
            rc.font.color.rgb = WARN_AMBER if kind == "residue" else DARK_BLUE
            rn = cn.paragraphs[0].add_run(name)
            rn.font.size = Pt(8.5)
            rn.font.color.rgb = DARK_GREY
            set_cell_bg(cc, shade)
            set_cell_bg(cn, shade)
        if len(pair) == 1:                     # odd count — blank the last pair
            for c in (2, 3):
                set_cell_bg(legend.rows[r].cells[c], "FFFFFF")
    section_break(doc)

    add_para(doc,
        "CAT is the honest residue: requirements no source in the research "
        "named, present because they are ordinary practice for a networked "
        "service that handles secrets. It never sits beside another code — the "
        "generator refuses to build if it does — so the CAT column counts "
        "exactly the requirements nothing evidenced, and those rows are marked "
        "in amber rather than left to hide among the others.",
        italic=True, size=9, colour=DARK_GREY)
    section_break(doc)
    add_para(doc,
        "Marks are given strictly, and that costs blank cells. A vendor whose "
        "documentation describes audit logging has evidenced an audit trail, "
        "not a tamper-evident one, so it leaves the tamper-evidence row blank. "
        "Read a blank as \u201cnot seen in what was read\u201d, never as "
        "\u201cabsent from the product\u201d.",
        italic=True, size=9, colour=DARK_GREY)
    doc.add_page_break()

    # ── the matrix, one table per domain ─────────────────────────────────
    per_source = {c: 0 for c in COLS}
    cat_only = []

    for domain in matrix.DOMAINS:
        add_heading(doc, domain["name"], 2, DARK_BLUE)
        tbl = doc.add_table(rows=1, cols=len(COLS) + 2)
        tbl.style = "Table Grid"
        matrix_header(tbl)
        for i, (feature, _expect, status, _route) in enumerate(domain["rows"]):
            marks = PROVENANCE[feature]
            for m in marks:
                per_source[m] += 1
            if marks == ["CAT"]:
                cat_only.append(feature)
            matrix_row(tbl, feature, marks, status, "EEF4FA" if i % 2 == 0 else "FFFFFF")
        widths(tbl)
        section_break(doc)

    doc.add_page_break()

    # ── what the columns add up to ───────────────────────────────────────
    add_heading(doc, "What the columns add up to", 1, DARK_BLUE)
    add_para(doc,
        "Read this as a description of the research, not of the market. A source "
        "with few marks is one whose documentation I read less of, not a product "
        "with fewer features.", size=9.5, italic=True)
    section_break(doc)

    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Table Grid"
    for i, text in enumerate(["Source", "Requirements evidenced",
                              f"Share of {TOTAL}", "Kind"]):
        cell = tbl.rows[0].cells[i]
        cell.text = text
        set_cell_bg(cell, "1A3A5C")
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(9)
    kind_label = {"vendor": "Vendor documentation", "standard": "Standard",
                  "residue": "Category baseline"}
    for i, (code, name, kind) in enumerate(SOURCES):
        row = tbl.add_row()
        vals = [f"{code} — {name}", str(per_source[code]),
                f"{per_source[code] * 100 // TOTAL} %", kind_label[kind]]
        for j, v in enumerate(vals):
            cell = row.cells[j]
            cell.text = v
            set_cell_bg(cell, "FBF3E3" if kind == "residue" else
                        ("EEF4FA" if i % 2 == 0 else "FFFFFF"))
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)
        row.cells[0].width = Inches(5.2)
        for j, w in enumerate([5.2, 1.8, 1.4, 2.19]):
            row.cells[j].width = Inches(w)
    section_break(doc)

    add_para(doc, f"Requirements carried only by the category baseline: "
                  f"{len(cat_only)} of {TOTAL}", bold=True, colour=WARN_AMBER)
    add_para(doc,
        "These are the rows to challenge first if the requirement list is ever "
        "disputed. Each is defensible as ordinary practice, but none was named "
        "by a vendor or a standard in the research behind the matrix.", size=9.5)
    for f in cat_only:
        add_bullet(doc, f, size=9)
    section_break(doc)

    add_para(doc, "What would make this stronger", bold=True, colour=MID_BLUE)
    for item in [
        "A hands-on evaluation of one or two products, which would turn this "
        "provenance table into a genuine capability comparison.",
        "The vendors' administration guides rather than their marketing pages — "
        "the former say what a product does, the latter what it advertises.",
        "An analyst grid, which would add an independent view of the category "
        "rather than each vendor's own.",
        "Direct access to four of the vendors' documentation sites. Bloombase, "
        "Cosmian, Securosys and HashiCorp were researched through indexed "
        "summaries of their pages because the network policy where this "
        "document is built blocks all four domains — so their columns are "
        "thinner than the reading would otherwise support.",
    ]:
        add_bullet(doc, item, size=9.5)

    section_break(doc)
    f = doc.add_paragraph()
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = f.add_run(f"KMIP on PKCS#11 — requirement provenance — {TODAY}")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    out = "KMIP_PKCS11_Feature_Provenance.docx"
    doc.save(out)
    print(f"Saved: {out}")
    print(f"  {TOTAL} requirements · {len(cat_only)} carried by the category "
          f"baseline alone")
    for code, _n, _k in SOURCES:
        print(f"  {code}: {per_source[code]}")


if __name__ == "__main__":
    build()
